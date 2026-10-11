#!/usr/bin/env python3
"""deck_qa — text-level QA for decks built with deck_lib, plus slide rendering.

    python deck_qa.py decks/<deck>.pptx                 # text checks
    python deck_qa.py decks/<deck>.pptx --render out/   # also render every slide to PNG

Text checks are necessary, not sufficient: the skills require inspecting every rendered slide.
Rendering needs LibreOffice (soffice) and poppler (pdftoppm) on PATH.
Exit code 1 when any FAIL is found.
"""
import argparse, os, re, subprocess, sys
from pptx import Presentation
from pptx.oxml.ns import qn

PITCH = ["unlock", "empower", "seamless", "game-changing", "industry-leading",
         "revolutionary", "supercharge"]
SOURCING = [r"\blearn (says|lists|states|notes|shows)\b", r"\baccording to learn\b"]
AMBIGUOUS_ACTORS = [r"\bagent team\b", r"\bAI team\b", r"\bmodel team\b"]
SECOND_PERSON = r"\b(you|your|yours|yourself)\b"
BADGES = r"^\s*(GA|Preview|PREVIEW|Public preview|Generally available)\s*$"
VERIFY = r"(Presenter note|\bverify\b[^.]*before presenting|\[verify|\bTODO\b|\bTBD\b)"
TITLE_Y_EMU = int(0.51 * 914400)
MIN_PT = 7.5
NOTES_MIN, NOTES_MAX = 120, 200


def shapes_flat(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from shapes_flat(sh.shapes)
        else:
            yield sh


def slide_texts(slide):
    for sh in shapes_flat(slide.shapes):
        if sh.has_text_frame:
            yield sh, sh.text_frame.text
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    yield sh, c.text


def find_title(slide):
    best = None
    for sh in shapes_flat(slide.shapes):
        if sh.has_text_frame and sh.top is not None and abs(sh.top - TITLE_Y_EMU) < 60000:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size.pt >= 24:
                        best = sh.text_frame.text.strip()
    return best


def looks_like_sentence(title):
    t = title.strip()
    if t.endswith((".", "?", "!", ":")):
        return True
    if len(t.split()) > 9:
        return True
    return bool(re.search(r"\b(is|are|was|were|does|do|can|will|should|must|lets|makes|runs|turns)\b", t, re.I))


def check(path):
    prs = Presentation(path)
    fails, warns = [], []
    for i, s in enumerate(prs.slides, 1):
        body = []
        small, has_star, has_foot = set(), False, False
        for sh, t in slide_texts(s):
            body.append(t)
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt < MIN_PT and r.text.strip():
                            small.add(r.font.size.pt)
            is_foot = bool(re.match(r"^\s*\*\s*preview", t, re.I))
            if "*" in t and not is_foot:
                has_star = True
            if is_foot:
                has_foot = True
            if re.search(BADGES, t):
                warns.append(f"slide {i}: status text '{t.strip()}' — allowed only in a Status table; elsewhere use an asterisk + footnote")
        text = "\n".join(body)
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        n_words = len(notes.split())
        if small:
            fails.append(f"slide {i}: font below {MIN_PT}pt: {sorted(small)}")
        if re.search(SECOND_PERSON, text, re.I):
            fails.append(f"slide {i}: second person (you/your) on the slide")
        for w in PITCH:
            if w in text.lower() or w in notes.lower():
                fails.append(f"slide {i}: pitch word '{w}'")
        for pat in SOURCING:
            if re.search(pat, text, re.I) or re.search(pat, notes, re.I):
                fails.append(f"slide {i}: '{pat}' — write 'as per official documentation' or drop the attribution")
        for pat in AMBIGUOUS_ACTORS:
            if re.search(pat, text, re.I) or re.search(pat, notes, re.I):
                warns.append(f"slide {i}: ambiguous actor ('{pat}') — name who: development team, platform team, operators")
        if re.search(VERIFY, text):
            fails.append(f"slide {i}: verify/TODO marker on the slide — move it to the speaker notes")
        if has_star and not has_foot:
            fails.append(f"slide {i}: asterisk without the preview footnote")
        if not notes.strip():
            fails.append(f"slide {i}: no speaker notes")
        elif not (NOTES_MIN <= n_words <= NOTES_MAX):
            warns.append(f"slide {i}: notes {n_words} words (target {NOTES_MIN}-{NOTES_MAX})")
        title = find_title(s)
        if i > 1 and title and looks_like_sentence(title):
            warns.append(f"slide {i}: title may be a sentence — use a short noun phrase: '{title}'")
        for sh in shapes_flat(s.shapes):
            sppr = sh._element.find(qn("p:spPr"))
            no_fx = sppr is not None and sppr.find(qn("a:effectLst")) is not None
            if sh._element.find(qn("p:style")) is not None and not no_fx:
                warns.append(f"slide {i}: shape '{sh.name}' keeps a theme style (possible shadow)")
                break
    return len(prs.slides), fails, warns


def render(path, outdir, dpi=110):
    os.makedirs(outdir, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), pdf, os.path.join(outdir, "slide")], check=True)
    return sorted(f for f in os.listdir(outdir) if f.endswith(".png"))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--render", metavar="DIR")
    a = ap.parse_args()
    n, fails, warns = check(a.deck)
    print(f"{a.deck}: {n} slides, {len(fails)} FAIL, {len(warns)} WARN")
    for f in fails:
        print("FAIL", f)
    for w in warns:
        print("WARN", w)
    if a.render:
        pngs = render(a.deck, a.render)
        print(f"rendered {len(pngs)} slide images to {a.render} — inspect every one")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
