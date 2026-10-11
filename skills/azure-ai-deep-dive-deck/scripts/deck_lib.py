"""deck_lib — the shared visual system for azure-ai-deep-dive-deck and foundry-whats-new-deck.

Implements Section 2 of azure-ai-deep-dive-deck (canonical) on python-pptx >= 1.0:
16:9 13.333 x 7.5 in, Segoe UI, fixed kicker/title/subtitle/footer positions, the navy/teal
palette, semantic diagram fills with a legend, preview asterisk + footnote, no shadows.

Usage
-----
    from deck_lib import *
    set_deck_footer("Hosted agents in Microsoft Foundry", "October 2026")
    prs = new_deck()
    title_slide(prs, "Hosted agents in Microsoft Foundry", "Subtitle ...",
                facts=[("Audience", "..."), ("Scope", "..."), ("Sources verified", "...")],
                presenter="<Name> · <Role>", date_label="October 2026")
    s = blank(prs)
    chrome(s, "01 · Overview", "Hosted agent anatomy", "One plain subtitle",
           "Source: learn.microsoft.com/azure/foundry/...", 2)
    notes(s, "120-200 words ...")
    prs.save("decks/hosted-agents-deep-dive-2026-10.pptx")

The presenter line is never hard-coded: pass the running user's name and role.
Change the palette or geometry here only together with Section 2 of both deck skills.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

# ---- palette ----
NAVY_DK = RGBColor(0x0B, 0x12, 0x29)
NAVY    = RGBColor(0x10, 0x1A, 0x3C)
TEAL    = RGBColor(0x00, 0xB3, 0x9F)
TEAL_DK = RGBColor(0x00, 0x81, 0x73)
CARD    = RGBColor(0xF6, 0xF8, 0xFD)
BORDER  = RGBColor(0xE1, 0xE7, 0xF4)
BAND    = RGBColor(0xEB, 0xEF, 0xF7)
INK     = RGBColor(0x14, 0x1B, 0x33)
MUTED   = RGBColor(0x44, 0x49, 0x5B)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LIGHTB  = RGBColor(0xB9, 0xC2, 0xD6)

FONT = "Segoe UI"

SW, SH = 13.333, 7.5
MX = 0.5
CW = SW - 2 * MX          # 12.333
FOOT_Y = 6.34
BODY_TOP = 1.45
BODY_BOT = 6.18

DECK_FOOTER = "<Topic> \u00b7 <Month Year>"


def set_deck_footer(topic, month_year):
    """Set the right-hand footer text used by chrome(): '<Topic> · <Month Year> · <n>'."""
    global DECK_FOOTER
    DECK_FOOTER = f"{topic} \u00b7 {month_year}"
    return DECK_FOOTER


def new_deck(title=None):
    """Blank 16:9 deck. Clears python-pptx's template metadata (author, 2013 dates)."""
    import datetime
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    cp = prs.core_properties
    now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
    cp.author = ""
    cp.last_modified_by = ""
    cp.created = now
    cp.modified = now
    cp.revision = 1
    if title:
        cp.title = title
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------- primitives ----------

def _fmt_run(r, size, color, bold=False, italic=False, font=FONT, spacing=None):
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    rPr = r._r.get_or_add_rPr()
    rPr.set('dirty', '0')
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        for e in rPr.findall(qn(tag)):
            rPr.remove(e)
    for tag in ('a:latin', 'a:cs'):
        e = rPr.makeelement(qn(tag), {'typeface': font})
        rPr.append(e)
    if spacing is not None:
        rPr.set('spc', str(int(spacing * 100)))


def textbox(slide, x, y, cx, cy, text, size=11, color=INK, bold=False,
            align='l', anchor='t', spacing=None, line=1.0, italic=False,
            space_after=0, wrap=True, shrink=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(cx), Inches(cy))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE,
                          'b': MSO_ANCHOR.BOTTOM}[anchor]
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER,
                       'r': PP_ALIGN.RIGHT, 'j': PP_ALIGN.JUSTIFY}[align]
        p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after)
        if isinstance(ln, tuple):
            for seg in ln:
                txt, opts = seg if isinstance(seg, tuple) else (seg, {})
                r = p.add_run(); r.text = txt
                _fmt_run(r, opts.get('size', size), opts.get('color', color),
                         opts.get('bold', bold), opts.get('italic', italic),
                         spacing=opts.get('spacing', spacing))
        else:
            r = p.add_run(); r.text = ln
            _fmt_run(r, size, color, bold, italic, spacing=spacing)
    return tb


def rect(slide, x, y, cx, cy, fill=None, line=None, lw=0.75, dash=None,
         preset=MSO_SHAPE.RECTANGLE, adj=None):
    sh = slide.shapes.add_shape(preset, Inches(x), Inches(y), Inches(cx), Inches(cy))
    if adj is not None and preset == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    # strip the theme shape style: it reintroduces the effect (shadow) reference
    sp = sh._element
    for st in sp.findall(qn('p:style')):
        sp.remove(st)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln()
            for e in ln.findall(qn('a:prstDash')):
                ln.remove(e)
            d = ln.makeelement(qn('a:prstDash'), {'val': dash})
            ln.append(d)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    sh.text_frame.text = ""
    return sh


def shape_text(sh, text, size=9, color=INK, bold=False, align='c', anchor='m',
               line=1.0, pad=0.06):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE,
                          'b': MSO_ANCHOR.BOTTOM}[anchor]
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER,
                       'r': PP_ALIGN.RIGHT}[align]
        p.line_spacing = line
        if isinstance(ln, tuple):
            for seg in ln:
                txt, opts = seg
                r = p.add_run(); r.text = txt
                _fmt_run(r, opts.get('size', size), opts.get('color', color),
                         opts.get('bold', bold), False)
        else:
            r = p.add_run(); r.text = ln
            _fmt_run(r, size, color, bold, False)
    return sh


def line(slide, x1, y1, x2, y2, color=BORDER, lw=0.75, dash=None, arrow=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(lw)
    ln = cn.line._get_or_add_ln()
    if dash:
        for e in ln.findall(qn('a:prstDash')):
            ln.remove(e)
        ln.append(ln.makeelement(qn('a:prstDash'), {'val': dash}))
    if arrow:
        for e in ln.findall(qn('a:tailEnd')):
            ln.remove(e)
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    for st in cn._element.findall(qn('p:style')):
        cn._element.remove(st)
    return cn


def elbow(slide, pts, color=MUTED, lw=1.0, dash=None, arrow=True):
    """Orthogonal polyline through pts [(x,y),...]; arrow on the last segment."""
    segs = []
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        segs.append(line(slide, x1, y1, x2, y2, color, lw, dash,
                         arrow and i == len(pts) - 2))
    return segs


def step(slide, x, y, n, d=0.23, fill=TEAL, color=WHITE, size=8.5):
    """Numbered circle placed ON a connector (x,y = centre)."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2),
                                Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = WHITE; sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    shape_text(sh, str(n), size=size, color=color, bold=True, pad=0)
    return sh


# ---------- slide furniture ----------

def chrome(slide, kicker, title, subtitle, source, page, footer=None):
    footer = footer or DECK_FOOTER
    textbox(slide, MX, 0.24, 8.0, 0.2, kicker.upper(), size=8, color=TEAL_DK,
            bold=True, spacing=1.4)
    textbox(slide, MX, 0.51, 12.0, 0.45, title, size=26, color=NAVY, bold=True)
    if subtitle:
        textbox(slide, MX, 1.02, 12.0, 0.3, subtitle, size=12.5, color=MUTED)
    if source:
        textbox(slide, MX, FOOT_Y, 7.6, 0.2, source, size=9, color=MUTED)
    textbox(slide, 9.2, FOOT_Y, SW - MX - 9.2, 0.2,
            f"{footer} \u00b7 {page}", size=9, color=MUTED, align='r')


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def card(slide, x, y, cx, cy, title, body, accent=TEAL, bar='top',
         tsize=10.5, bsize=8.5, tcolor=NAVY, bcolor=MUTED, fill=CARD,
         border=BORDER, pad=0.16):
    sh = rect(slide, x, y, cx, cy, fill=fill, line=border, adj=0.06)
    if bar == 'top':
        b = rect(slide, x, y, cx, 0.045, fill=accent, preset=MSO_SHAPE.RECTANGLE)
    elif bar == 'left':
        b = rect(slide, x, y, 0.045, cy, fill=accent, preset=MSO_SHAPE.RECTANGLE)
    ty = y + (0.13 if bar == 'top' else 0.11)
    textbox(slide, x + pad, ty, cx - 2 * pad, 0.22, title, size=tsize,
            color=tcolor, bold=True, line=0.95)
    if body:
        textbox(slide, x + pad, ty + 0.265, cx - 2 * pad, cy - (ty - y) - 0.32,
                body, size=bsize, color=bcolor, line=1.12, space_after=2.5)
    return sh


def legend(slide, items, x=MX, y=6.03, size=7.5, gap=0.14):
    """items: [(kind,label)] kind in product|capability|channel|governance|boundary|solid|dashed"""
    cx = x
    for kind, label in items:
        if kind == 'product':
            rect(slide, cx, y + 0.035, 0.16, 0.1, fill=NAVY, preset=MSO_SHAPE.RECTANGLE)
        elif kind == 'capability':
            rect(slide, cx, y + 0.035, 0.16, 0.1, fill=WHITE, line=BORDER, lw=0.75,
                 preset=MSO_SHAPE.RECTANGLE)
        elif kind == 'channel':
            rect(slide, cx, y + 0.035, 0.16, 0.1, fill=BAND, preset=MSO_SHAPE.RECTANGLE)
        elif kind == 'governance':
            rect(slide, cx, y + 0.035, 0.16, 0.1, fill=NAVY_DK, preset=MSO_SHAPE.RECTANGLE)
        elif kind == 'boundary':
            rect(slide, cx, y + 0.02, 0.16, 0.13, fill=None, line=MUTED, lw=0.75,
                 dash='dash', preset=MSO_SHAPE.RECTANGLE)
        elif kind == 'solid':
            line(slide, cx, y + 0.085, cx + 0.18, y + 0.085, color=MUTED, lw=1.1)
        elif kind == 'dashed':
            line(slide, cx, y + 0.085, cx + 0.18, y + 0.085, color=MUTED, lw=1.1, dash='dash')
        elif kind == 'teal':
            rect(slide, cx, y + 0.035, 0.16, 0.1, fill=TEAL, preset=MSO_SHAPE.RECTANGLE)
        cx += 0.24
        tb = textbox(slide, cx, y + 0.005, 3.2, 0.16, label, size=size, color=MUTED, wrap=False)
        cx += 0.14 + len(label) * size * 0.0080


def footnote(slide, text, y=6.15, x=MX, size=7.5):
    textbox(slide, x, y, 8.0, 0.16, text, size=size, color=MUTED, italic=True)


def rowtable(slide, x, y, cx, headers, rows, colw, rowh=0.44, hsize=8.5,
             bsize=8.5, head_fill=NAVY, zebra=True, first_bold=True):
    """Parallel-row comparison table: same question asked of every column.
    colw: column widths as FRACTIONS of cx (sum to 1.0). rows: [cells] or ([cells], height).
    Every row must answer the table's own question (Slide wording rule 6)."""
    cw = [cx * w for w in colw]
    hx = x
    hd = rect(slide, x, y, cx, 0.34, fill=head_fill, line=None, adj=0.10)
    for i, h in enumerate(headers):
        textbox(slide, hx + 0.14, y + 0.085, cw[i] - 0.2, 0.2, h, size=hsize,
                color=WHITE, bold=True, wrap=False)
        hx += cw[i]
    yy = y + 0.34
    for ri, row in enumerate(rows):
        if isinstance(row, tuple) and len(row) == 2 and isinstance(row[1], (int, float)):
            cells, h = row
        else:
            cells, h = row, rowh
        if zebra and ri % 2 == 0:
            rect(slide, x, yy, cx, h, fill=CARD, line=None, preset=MSO_SHAPE.RECTANGLE)
        line(slide, x, yy + h, x + cx, yy + h, color=BORDER, lw=0.75)
        hx = x
        for ci, c in enumerate(cells):
            textbox(slide, hx + 0.14, yy + 0.055, cw[ci] - 0.24, h - 0.09, c,
                    size=bsize, color=NAVY if (ci == 0 and first_bold) else MUTED,
                    bold=(ci == 0 and first_bold), anchor='m', line=1.12)
            hx += cw[ci]
        yy += h
    for i in range(1, len(headers)):
        vx = x + sum(cw[:i])
        line(slide, vx, y + 0.34, vx, yy, color=BORDER, lw=0.75)
    return yy


def chip(slide, x, y, cx, cy, text, fill=BAND, color=NAVY, size=8, bold=False,
         border=None):
    sh = rect(slide, x, y, cx, cy, fill=fill, line=border, adj=0.30)
    shape_text(sh, text, size=size, color=color, bold=bold)
    return sh


def kv(slide, x, y, cx, pairs, kw=0.30, size=8.5, rowh=0.235, rule=True):
    """Key-value rows."""
    yy = y
    for k, v in pairs:
        textbox(slide, x, yy, cx * kw - 0.1, rowh, k, size=size, color=NAVY,
                bold=True, anchor='m', line=1.05)
        textbox(slide, x + cx * kw, yy, cx * (1 - kw), rowh, v, size=size,
                color=MUTED, anchor='m', line=1.05)
        yy += rowh
        if rule:
            line(slide, x, yy - 0.02, x + cx, yy - 0.02, color=BORDER, lw=0.75)
    return yy


MONO = "Consolas"


def code(slide, x, y, cx, cy, lines, size=8, title=None, pad=0.24):
    rect(slide, x, y, cx, cy, fill=NAVY_DK, line=None)
    ty = y + 0.16
    if title:
        textbox(slide, x + pad, ty, cx - 2 * pad, 0.18, title.upper(), size=7.5,
                color=TEAL, bold=True, spacing=1.1)
        ty += 0.26
    tb = textbox(slide, x + pad, ty, cx - 2 * pad, cy - (ty - y) - 0.14, lines,
                 size=size, color=RGBColor(0xD6, 0xDE, 0xEE), line=1.28, wrap=False)
    for para in tb.text_frame.paragraphs:
        for r in para.runs:
            r.font.name = MONO
            rPr = r._r.get_or_add_rPr()
            for tag in ('a:latin', 'a:cs'):
                for e in rPr.findall(qn(tag)):
                    e.set('typeface', MONO)
    return tb


def band(slide, x, y, cx, cy, text, fill=BAND, color=NAVY, size=9, accent=TEAL,
         pad=0.26, line=1.2):
    rect(slide, x, y, cx, cy, fill=fill, line=None)
    if accent is not None:
        rect(slide, x, y, 0.045, cy, fill=accent)
    textbox(slide, x + pad, y + 0.11, cx - pad - 0.2, cy - 0.20, text, size=size,
            color=color, line=line)


def sectionlabel(slide, x, y, text, cx=7.0):
    textbox(slide, x, y, cx, 0.2, text.upper(), size=8, color=TEAL_DK, bold=True,
            spacing=1.2)


# ---------- diagram and slide helpers ----------

PREVIEW_FOOTNOTE = "* preview \u2014 no service level agreement."


def pfoot(s, y=6.05):
    """Right-aligned preview footnote; required on every slide that uses an asterisk."""
    textbox(s, 8.6, y, SW - MX - 8.6, 0.16, PREVIEW_FOOTNOTE,
            size=7.5, color=MUTED, italic=True, align='r')


def arrow(s, x1, y1, x2, y2, color=MUTED, lw=1.1, dash=None):
    """Straight connector with arrow head. Solid = data path; dash='dash' = identity/control/telemetry."""
    return line(s, x1, y1, x2, y2, color=color, lw=lw, dash=dash, arrow=True)


_KINDS = {'cap': (WHITE, BORDER, NAVY, MUTED), 'chan': (BAND, None, NAVY, MUTED),
          'prod': (NAVY, None, WHITE, LIGHTB), 'gov': (NAVY_DK, None, WHITE, LIGHTB)}


def box(s, x, y, w, h, text, kind='cap', size=8.5, bold=False, align='c'):
    """Diagram box. kind: prod (product described) | cap (capability/resource) |
    chan (channel, caller, customer-owned) | gov (governance/control layer)."""
    fill, ln, col, _ = _KINDS[kind]
    sh = rect(s, x, y, w, h, fill=fill, line=ln)
    shape_text(sh, text, size=size, color=col, bold=bold, align=align, line=1.12)
    return sh


def labelled(s, x, y, w, h, title, sub, kind='cap', tsize=8.5, ssize=7.5):
    """Diagram box with a bold title line and a smaller second line."""
    fill, ln, tc, sc = _KINDS[kind]
    sh = rect(s, x, y, w, h, fill=fill, line=ln)
    shape_text(sh, [((title, {'bold': True, 'color': tc, 'size': tsize}),),
                    ((sub, {'color': sc, 'size': ssize}),)], line=1.1)
    return sh


def title_slide(prs, title, subtitle, facts, presenter, date_label,
                eyebrow="MICROSOFT FOUNDRY  \u00b7  TECHNICAL DEEP DIVE", title_size=None):
    """Dark title slide. facts: up to three (label, value) pairs. presenter: the running
    user's name and role, read from their profile — never a hard-coded name."""
    s = blank(prs)
    rect(s, 0, 0, SW, SH, fill=NAVY_DK)
    rect(s, 0, 0, 0.09, SH, fill=TEAL)
    textbox(s, 1.0, 2.00, 10.0, 0.22, eyebrow, size=9, color=TEAL, bold=True, spacing=1.6)
    size = title_size or (42 if len(title) <= 26 else 36)
    textbox(s, 1.0, 2.41, 11.8, 0.72, title, size=size, color=WHITE, bold=True, wrap=False)
    textbox(s, 1.0, 3.24, 10.8, 0.62, subtitle, size=14.5, color=LIGHTB, line=1.32)
    line(s, 1.0, 4.26, 3.2, 4.26, color=TEAL, lw=2.0)
    for i, (k, v) in enumerate(facts[:3]):
        x = 1.0 + i * 3.7
        textbox(s, x, 4.56, 3.5, 0.18, k.upper(), size=7.5, color=TEAL, bold=True, spacing=1.2)
        textbox(s, x, 4.80, 3.5, 0.52, v, size=9.5, color=LIGHTB, line=1.28)
    if presenter:
        textbox(s, 1.0, 6.45, 8.0, 0.2, presenter, size=9.5, color=LIGHTB)
    textbox(s, 9.2, 6.45, SW - 1.0 - 9.2, 0.2, date_label, size=9.5, color=LIGHTB, align='r')
    return s
